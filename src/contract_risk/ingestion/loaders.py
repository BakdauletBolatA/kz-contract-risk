"""Загрузка TXT, DOCX и PDF в нормализованный текст.

DOCX. python-docx отдаёт текст абзаца без номера, если номер — автонумерация
Word, а в договорах это почти всегда так. Без восстановления номеров
«5.2. Арендатор уплачивает…» превращается в «Арендатор уплачивает…», и
сегментатору не на что опереться. Поэтому номера пересчитываются по
numbering.xml: счётчики по уровням, формат уровня (1, а, i…), шаблон
вида «%1.%2.». Таблицы обходятся в порядке документа; двуязычный договор
(две колонки kk | ru) сводится к колонке нужного языка.

PDF. Текстовый слой через pdfplumber, колонтитулы убираются по повторам.
Если текста почти нет — это скан: документ помечается `needs_ocr`, и без
включённого OCR анализ честно не выполняется.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from contract_risk.ingestion.language import detect_language
from contract_risk.ingestion.normalize import normalize_text, strip_repeated_lines
from contract_risk.ingestion.ocr import OcrEngine
from contract_risk.schemas import Language

TEXT_EXTENSIONS = {".txt"}
DOCX_EXTENSIONS = {".docx"}
PDF_EXTENSIONS = {".pdf"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | DOCX_EXTENSIONS | PDF_EXTENSIONS | IMAGE_EXTENSIONS

# Меньше этого числа символов на страницу в среднем — считаем PDF сканом.
SCAN_CHARS_PER_PAGE = 50


class UnsupportedFormatError(ValueError):
    pass


@dataclass
class RawDocument:
    text: str
    source_format: str
    page_starts: list[int] | None = None
    needs_ocr: bool = False
    ocr_used: bool = False
    warnings: list[str] = field(default_factory=list)


def load_bytes(
    data: bytes,
    filename: str,
    *,
    ocr: OcrEngine | None = None,
    prefer_language: Language = Language.RU,
) -> RawDocument:
    ext = Path(filename).suffix.lower()
    if ext in TEXT_EXTENSIONS:
        return _load_txt(data)
    if ext in DOCX_EXTENSIONS:
        return _load_docx(data, prefer_language)
    if ext in PDF_EXTENSIONS:
        return _load_pdf(data, ocr)
    if ext in IMAGE_EXTENSIONS:
        return _load_image(data, ocr)
    raise UnsupportedFormatError(
        f"формат {ext or '(без расширения)'} не поддерживается: "
        f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def load_path(path: Path, **kwargs) -> RawDocument:
    return load_bytes(Path(path).read_bytes(), Path(path).name, **kwargs)


# --- TXT ---------------------------------------------------------------------


def _load_txt(data: bytes) -> RawDocument:
    warnings = []
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1251", errors="replace")
        warnings.append("файл не в UTF-8: прочитан как Windows-1251")
    return RawDocument(normalize_text(text), "txt", warnings=warnings)


# --- DOCX --------------------------------------------------------------------

_RU_LOWER = "абвгдежзиклмнопрстуфхцчшщэюя"  # как russianLower в Word: без ё, й, ъ, ы, ь


def _roman(n: int) -> str:
    out = ""
    for value, sym in (
        (1000, "m"),
        (900, "cm"),
        (500, "d"),
        (400, "cd"),
        (100, "c"),
        (90, "xc"),
        (50, "l"),
        (40, "xl"),
        (10, "x"),
        (9, "ix"),
        (5, "v"),
        (4, "iv"),
        (1, "i"),
    ):
        while n >= value:
            out += sym
            n -= value
    return out


def _letters(n: int, alphabet: str) -> str:
    # Word после «z» пишет «aa», «bb»…: буква повторяется.
    return alphabet[(n - 1) % len(alphabet)] * ((n - 1) // len(alphabet) + 1)


def format_number(value: int, fmt: str) -> str:
    match fmt:
        case "lowerLetter":
            return _letters(value, "abcdefghijklmnopqrstuvwxyz")
        case "upperLetter":
            return _letters(value, "ABCDEFGHIJKLMNOPQRSTUVWXYZ")
        case "russianLower":
            return _letters(value, _RU_LOWER)
        case "russianUpper":
            return _letters(value, _RU_LOWER.upper())
        case "lowerRoman":
            return _roman(value)
        case "upperRoman":
            return _roman(value).upper()
        case "decimalZero":
            return f"{value:02d}"
        case _:
            return str(value)


@dataclass
class _Level:
    start: int = 1
    fmt: str = "decimal"
    text: str = ""
    is_legal: bool = False


class WordNumbering:
    """Пересчёт автонумерации Word по numbering.xml.

    Счётчики ведутся по abstractNum: в Word списки с разными numId, но общим
    abstractNum продолжают друг друга, а «начать заново» — это отдельный
    w:num со startOverride.
    """

    def __init__(self, numbering_element=None) -> None:  # noqa: ANN001 — lxml element
        from docx.oxml.ns import qn

        self._qn = qn
        self.levels: dict[str, dict[int, _Level]] = {}
        self.nums: dict[str, tuple[str, dict[int, int]]] = {}
        self._counters: dict[str, list[int | None]] = {}
        self._started: set[str] = set()
        if numbering_element is None:
            return
        for abstract in numbering_element.findall(qn("w:abstractNum")):
            levels = {}
            for lvl in abstract.findall(qn("w:lvl")):
                levels[int(lvl.get(qn("w:ilvl")))] = _Level(
                    start=int(self._val(lvl, "w:start", "1")),
                    fmt=self._val(lvl, "w:numFmt", "decimal"),
                    text=self._val(lvl, "w:lvlText", ""),
                    is_legal=lvl.find(qn("w:isLgl")) is not None,
                )
            self.levels[abstract.get(qn("w:abstractNumId"))] = levels
        for num in numbering_element.findall(qn("w:num")):
            overrides = {}
            for ov in num.findall(qn("w:lvlOverride")):
                start = ov.find(qn("w:startOverride"))
                if start is not None:
                    overrides[int(ov.get(qn("w:ilvl")))] = int(start.get(qn("w:val")))
            self.nums[num.get(qn("w:numId"))] = (self._val(num, "w:abstractNumId", ""), overrides)

    def _val(self, element, tag: str, default: str) -> str:  # noqa: ANN001
        child = element.find(self._qn(tag))
        if child is None:
            return default
        return child.get(self._qn("w:val"), default)

    def label(self, num_id: str, ilvl: int) -> str | None:
        if num_id == "0" or num_id not in self.nums:
            return None
        abstract_id, overrides = self.nums[num_id]
        levels = self.levels.get(abstract_id, {})
        if ilvl not in levels:
            return None
        counters = self._counters.setdefault(abstract_id, [None] * 9)
        if overrides and num_id not in self._started:
            # Новый экземпляр списка с «начать заново».
            counters[:] = [None] * 9
        self._started.add(num_id)

        level = levels[ilvl]
        current = counters[ilvl]
        counters[ilvl] = overrides.get(ilvl, level.start) if current is None else current + 1
        for deeper in range(ilvl + 1, 9):
            counters[deeper] = None

        if level.fmt == "bullet":
            return "•"
        if level.fmt == "none":
            return ""

        def render(m: re.Match) -> str:
            k = int(m.group(1)) - 1
            ref = levels.get(k, _Level())
            value = counters[k] if counters[k] is not None else overrides.get(k, ref.start)
            return format_number(value, "decimal" if level.is_legal else ref.fmt)

        return re.sub(r"%(\d)", render, level.text)


def _num_pr(paragraph) -> tuple[str, int] | None:  # noqa: ANN001 — docx Paragraph
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.numPr is not None and p_pr.numPr.numId is not None:
        ilvl = p_pr.numPr.ilvl.val if p_pr.numPr.ilvl is not None else 0
        return str(p_pr.numPr.numId.val), int(ilvl)
    style = paragraph.style
    while style is not None:
        s_pr = style.element.pPr
        if s_pr is not None and s_pr.numPr is not None and s_pr.numPr.numId is not None:
            ilvl = s_pr.numPr.ilvl.val if s_pr.numPr.ilvl is not None else 0
            return str(s_pr.numPr.numId.val), int(ilvl)
        style = style.base_style
    return None


def _paragraph_text(paragraph, numbering: WordNumbering) -> str:  # noqa: ANN001
    text = paragraph.text.strip()
    num = _num_pr(paragraph)
    if num is None or not text:
        return text
    label = numbering.label(*num)
    return f"{label} {text}" if label else text


def _load_docx(data: bytes, prefer_language: Language) -> RawDocument:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    try:
        numbering_element = document.part.numbering_part.element
    except (KeyError, NotImplementedError):
        numbering_element = None
    numbering = WordNumbering(numbering_element)
    warnings: list[str] = []
    lines: list[str] = []

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            lines.append(_paragraph_text(Paragraph(child, document), numbering))
        elif child.tag == qn("w:tbl"):
            lines.extend(_table_lines(Table(child, document), numbering, prefer_language, warnings))
    return RawDocument(normalize_text("\n".join(lines)), "docx", warnings=warnings)


def _table_lines(table, numbering, prefer_language, warnings) -> list[str]:  # noqa: ANN001
    rows = []
    for row in table.rows:
        seen, cells = set(), []
        for cell in row.cells:
            if id(cell._tc) not in seen:  # объединённые ячейки повторяются в row.cells
                seen.add(id(cell._tc))
                cells.append(cell)
        rows.append(cells)

    two_col = [r for r in rows if len(r) == 2]
    if rows and len(two_col) >= 0.6 * len(rows):
        columns = [
            "\n".join(p.text for r in two_col for p in r[i].paragraphs).strip() for i in (0, 1)
        ]
        langs = [detect_language(c) for c in columns]
        if all(columns) and langs[0] != langs[1]:
            keep = langs.index(prefer_language) if prefer_language in langs else 1
            warnings.append(
                f"двуязычный договор: анализируется колонка «{langs[keep].value}», "
                f"вторая колонка пропущена"
            )
            return [_paragraph_text(p, numbering) for r in two_col for p in r[keep].paragraphs]

    lines = []
    for cells in rows:
        texts = ["\n".join(_paragraph_text(p, numbering) for p in c.paragraphs) for c in cells]
        texts = [t.strip() for t in texts if t.strip()]
        if len(texts) == 1:
            lines.extend(texts[0].split("\n"))
        elif texts:
            lines.append(" | ".join(t.replace("\n", " ") for t in texts))
    return lines


# --- PDF и изображения -------------------------------------------------------


def _assemble_pages(pages: list[str]) -> tuple[str, list[int]]:
    pages = strip_repeated_lines([normalize_text(p) for p in pages])
    starts, parts, offset = [], [], 0
    for page in pages:
        page = page if page.endswith("\n") or not page else page + "\n"
        starts.append(offset)
        parts.append(page)
        offset += len(page)
    return "".join(parts), starts


def _load_pdf(data: bytes, ocr: OcrEngine | None) -> RawDocument:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = [page.extract_text(x_tolerance=1.5, y_tolerance=3) or "" for page in pdf.pages]

    warnings: list[str] = []
    joined = "".join(pages)
    broken_font = joined.count("(cid:") > 20
    scanned = not pages or sum(len(p.strip()) for p in pages) / len(pages) < SCAN_CHARS_PER_PAGE
    if broken_font:
        warnings.append("в PDF нет таблицы символов шрифта: текстовый слой нечитаем")
    if scanned or broken_font:
        if ocr is None:
            warnings.append(
                "документ похож на скан: текстового слоя нет, включите OCR или загрузите DOCX"
            )
            return RawDocument("", "pdf", needs_ocr=True, warnings=warnings)
        pages = ocr.pdf_pages(data)
        warnings.append(f"текст получен распознаванием ({ocr.name}): возможны ошибки в цитатах")
        text, starts = _assemble_pages(pages)
        return RawDocument(text, "pdf", starts, needs_ocr=True, ocr_used=True, warnings=warnings)

    text, starts = _assemble_pages(pages)
    return RawDocument(text, "pdf", starts, warnings=warnings)


def _load_image(data: bytes, ocr: OcrEngine | None) -> RawDocument:
    if ocr is None:
        return RawDocument(
            "",
            "pdf",
            needs_ocr=True,
            warnings=["изображение: для анализа нужен OCR, он не включён"],
        )
    text = normalize_text(ocr.image(data))
    return RawDocument(
        text,
        "pdf",
        [0],
        needs_ocr=True,
        ocr_used=True,
        warnings=[f"текст получен распознаванием ({ocr.name}): возможны ошибки в цитатах"],
    )
