"""Загрузка DOCX (автонумерация Word, таблицы), PDF, сканов и OCR."""

import io
import shutil

import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from contract_risk.ingestion import parse_bytes
from contract_risk.ingestion.loaders import (
    UnsupportedFormatError,
    WordNumbering,
    format_number,
    load_bytes,
)
from contract_risk.ingestion.ocr import TesseractOcr
from contract_risk.schemas import Language

ABSTRACT = f"""<w:abstractNum {nsdecls("w")} w:abstractNumId="90">
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>
  <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>
  <w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="russianLower"/><w:lvlText w:val="%3)"/></w:lvl>
</w:abstractNum>"""
NUM = f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'


def numbered_docx(paragraphs: list[tuple[int, str]]) -> bytes:
    document = docx.Document()
    numbering = document.part.numbering_part.element
    numbering.insert(0, parse_xml(ABSTRACT))
    numbering.append(parse_xml(NUM))
    document.add_paragraph("ДОГОВОР АРЕНДЫ")
    for ilvl, text in paragraphs:
        p = document.add_paragraph(text)
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = ilvl
        num_pr.get_or_add_numId().val = 90
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_word_autonumbering_is_restored():
    data = numbered_docx(
        [
            (0, "ПРЕДМЕТ ДОГОВОРА"),
            (1, "Арендодатель передает помещение."),
            (1, "Помещение используется под офис."),
            (0, "ОТВЕТСТВЕННОСТЬ СТОРОН"),
            (1, "Арендатор обязан:"),
            (2, "вносить плату;"),
            (2, "содержать помещение."),
        ]
    )
    raw = load_bytes(data, "lease.docx")
    assert raw.text.splitlines() == [
        "ДОГОВОР АРЕНДЫ",
        "1. ПРЕДМЕТ ДОГОВОРА",
        "1.1. Арендодатель передает помещение.",
        "1.2. Помещение используется под офис.",
        "2. ОТВЕТСТВЕННОСТЬ СТОРОН",
        "2.1. Арендатор обязан:",
        "а) вносить плату;",
        "б) содержать помещение.",
    ]
    doc = parse_bytes(data, "lease.docx")
    assert [c.id for c in doc.analysable_clauses] == ["1.1", "1.2", "2.1"]
    assert doc.clause("2.1").text.endswith("б) содержать помещение.")


def test_numbering_restart_via_start_override():
    xml = parse_xml(
        f"""<w:numbering {nsdecls("w")}>{ABSTRACT}{NUM}
        <w:num w:numId="91"><w:abstractNumId w:val="90"/>
          <w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>
        </w:numbering>"""
    )
    numbering = WordNumbering(xml)
    assert [numbering.label("90", 0) for _ in range(2)] == ["1.", "2."]
    assert numbering.label("91", 0) == "1."
    assert numbering.label("0", 0) is None


@pytest.mark.parametrize(
    ("value", "fmt", "expected"),
    [
        (3, "decimal", "3"),
        (2, "lowerLetter", "b"),
        (27, "lowerLetter", "aa"),
        (10, "russianLower", "к"),  # в Word нет «й»: после «и» идёт «к»
        (4, "upperRoman", "IV"),
        (7, "decimalZero", "07"),
    ],
)
def test_format_number(value, fmt, expected):
    assert format_number(value, fmt) == expected


def test_bilingual_table_keeps_preferred_language_column():
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "1.1. Жалға алушы ай сайын жалдау ақысын төлейді."
    table.cell(0, 1).text = "1.1. Арендатор ежемесячно вносит арендную плату."
    table.cell(1, 0).text = "1.2. Жалға беруші үй-жайды береді."
    table.cell(1, 1).text = "1.2. Арендодатель передает помещение."
    buf = io.BytesIO()
    document.save(buf)

    ru = parse_bytes(buf.getvalue(), "bi.docx")
    assert [c.id for c in ru.analysable_clauses] == ["1.1", "1.2"]
    assert ru.language == Language.RU
    assert any("двуязычный" in w for w in ru.warnings)

    kk = parse_bytes(buf.getvalue(), "bi.docx", prefer_language=Language.KK)
    assert kk.language == Language.KK
    assert kk.clause("1.1").text.startswith("Жалға алушы")


def test_requisites_table_is_flattened_row_by_row():
    document = docx.Document()
    document.add_paragraph("1. Предмет договора: поставка товара.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Поставщик: ТОО «А»"
    table.cell(0, 1).text = "Покупатель: ИП «Б»"
    buf = io.BytesIO()
    document.save(buf)
    raw = load_bytes(buf.getvalue(), "s.docx")
    assert "Поставщик: ТОО «А» | Покупатель: ИП «Б»" in raw.text


def test_txt_in_cp1251_is_read_with_warning():
    raw = load_bytes("1.1. Арендатор платит.".encode("cp1251"), "a.txt")
    assert raw.text.startswith("1.1. Арендатор")
    assert raw.warnings


def test_unsupported_extension():
    with pytest.raises(UnsupportedFormatError):
        load_bytes(b"", "contract.odt")


def _weasyprint():
    try:
        import weasyprint
    except (ImportError, OSError):
        pytest.skip("WeasyPrint недоступен — PDF для теста не собрать")
    return weasyprint


def test_pdf_text_layer_pages_and_headers():
    weasyprint = _weasyprint()
    pages = []
    for i in range(1, 4):
        pages.append(
            f"<div style='page-break-after: always'><p>ТОО Колонтитул</p>"
            f"<p>{i}.1. Пункт номер {i} договора аренды.</p></div>"
        )
    html = f"<html><body style='font-family: Arial'>{''.join(pages)}</body></html>"
    data = weasyprint.HTML(string=html).write_pdf()
    doc = parse_bytes(data, "c.pdf")
    assert doc.source_format == "pdf"
    assert "Колонтитул" not in doc.text
    assert [c.id for c in doc.analysable_clauses] == ["1.1", "2.1", "3.1"]
    assert [c.page for c in doc.analysable_clauses] == [1, 2, 3]


def test_pdf_without_text_layer_is_flagged_as_scan():
    weasyprint = _weasyprint()
    html = "<html><body><div style='width:100px;height:100px;background:#000'></div></body></html>"
    doc = parse_bytes(weasyprint.HTML(string=html).write_pdf(), "scan.pdf")
    assert doc.needs_ocr and not doc.ocr_used
    assert doc.clauses == []
    assert any("скан" in w for w in doc.warnings)


FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"


@pytest.mark.skipif(not TesseractOcr.available(), reason="нет tesseract/pytesseract")
def test_ocr_reads_rendered_cyrillic_image():
    from PIL import Image, ImageDraw, ImageFont

    if not shutil.os.path.exists(FONT):
        pytest.skip("нет шрифта с кириллицей для рендера")
    image = Image.new("RGB", (1400, 220), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(FONT, 44)
    draw.text((20, 30), "1.1. Арендатор уплачивает пеню.", fill="black", font=font)
    draw.text((20, 120), "1.2. Арендодатель передает помещение.", fill="black", font=font)
    buf = io.BytesIO()
    image.save(buf, format="PNG")

    doc = parse_bytes(buf.getvalue(), "scan.png", ocr=TesseractOcr())
    assert doc.ocr_used
    assert "Арендатор" in doc.text
    assert [c.id for c in doc.analysable_clauses] == ["1.1", "1.2"]
